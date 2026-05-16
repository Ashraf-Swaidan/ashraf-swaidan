from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).with_name("Ashraf-Swaidan-Resume-Portfolio-First.docx")

INK = RGBColor(22, 28, 36)
MUTED = RGBColor(88, 98, 110)
ACCENT = RGBColor(20, 86, 125)
SOFT = RGBColor(236, 241, 245)


def set_paragraph_border(paragraph, color="D4DCE4", size="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_hyperlink(paragraph, text, url, bold=False):
    rel_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)

    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    if bold:
        r_pr.append(OxmlElement("w:b"))
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "14567D")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "none")
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def run(paragraph, text, *, bold=False, italic=False, size=9.4, color=INK):
    r = paragraph.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    r.font.color.rgb = color
    return r


def section(doc, title):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run(p, title.upper(), bold=True, size=10.4, color=ACCENT)
    set_paragraph_border(p)


def text(doc, value, after=4, size=9.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.02
    run(p, value, size=size)
    return p


def bullet(doc, value, after=2):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    run(p, value, size=9.1)


def role(doc, title, meta):
    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.05), WD_TAB_ALIGNMENT.RIGHT)
    run(p, title, bold=True, size=10.1)
    run(p, f"\t{meta}", size=8.9, color=MUTED)


def project(doc, title, case_url, lines, live_url=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(1)
    run(p, title, bold=True, size=10.0)

    link = doc.add_paragraph()
    link.paragraph_format.space_after = Pt(1)
    run(link, "Portfolio case study: ", size=8.7, color=MUTED)
    add_hyperlink(link, case_url.replace("https://", ""), case_url, bold=True)
    if live_url:
        run(link, " | Live: ", size=8.7, color=MUTED)
        add_hyperlink(link, live_url.replace("https://", ""), live_url)

    for value in lines:
        bullet(doc, value, after=1.7)


def skill_line(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run(p, f"{label}: ", bold=True, size=9.0, color=ACCENT)
    run(p, value, size=9.0)


def configure_doc(doc):
    sec = doc.sections[0]
    sec.start_type = WD_SECTION.NEW_PAGE
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.52)
    sec.bottom_margin = Inches(0.5)
    sec.left_margin = Inches(0.58)
    sec.right_margin = Inches(0.58)

    for style_name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "List Bullet"]:
        style = doc.styles[style_name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
        style.font.color.rgb = INK
    doc.styles["Normal"].font.size = Pt(9.4)
    doc.styles["Title"].font.size = Pt(21)
    doc.styles["Title"].font.bold = True
    doc.styles["Subtitle"].font.size = Pt(10.6)
    doc.styles["Subtitle"].font.color.rgb = MUTED
    doc.styles["List Bullet"].font.size = Pt(9.1)


def build():
    doc = Document()
    configure_doc(doc)

    name = doc.add_paragraph(style="Title")
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(0)
    name.add_run("Ashraf Emad Swaidan")

    sub = doc.add_paragraph(style="Subtitle")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(2)
    sub.add_run("Product-minded Software Developer | Operational Systems & Internal Tools")

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(2)
    for i, part in enumerate(
        [
            ("Lebanon", None),
            ("+961 76 350 373", "tel:+96176350373"),
            ("ashraf.swaidan.13@gmail.com", "mailto:ashraf.swaidan.13@gmail.com"),
            ("GitHub", "https://github.com/Ashraf-Swaidan"),
            ("LinkedIn", "https://www.linkedin.com/in/ashraf-swaidan-9aaa612b1/"),
        ]
    ):
        if i:
            run(contact, " | ", size=8.9, color=MUTED)
        label, url = part
        if url:
            add_hyperlink(contact, label, url)
        else:
            run(contact, label, size=8.9, color=MUTED)

    portfolio = doc.add_paragraph()
    portfolio.alignment = WD_ALIGN_PARAGRAPH.CENTER
    portfolio.paragraph_format.space_after = Pt(5)
    run(portfolio, "Portfolio: ", bold=True, size=10.0, color=ACCENT)
    add_hyperlink(portfolio, "ashraf-swaidan.pages.dev", "https://ashraf-swaidan.pages.dev", bold=True)
    set_paragraph_border(portfolio, color="E3E8EE", size="6")

    section(doc, "Profile")
    text(
        doc,
        "Software developer and Computer Science graduate who builds systems for real business operations. My strongest work is best understood through the portfolio: production-style platforms that turn messy workflows into clear, usable products for sales, inventory, customers, finance, suppliers, tasks, analytics, and daily execution.",
        after=3,
    )

    section(doc, "Portfolio Highlights")
    project(
        doc,
        "Papion System - Business Operations Platform",
        "https://ashraf-swaidan.pages.dev/works/papion-system",
        [
            "Built a production operations system for a multi-branch event decoration and print/cut business.",
            "Unified sales, inventory, customers, suppliers, expenses, wallets, tasks, calendars, analytics, roles, permissions, and AI-assisted business Q&A.",
            "Designed workflows around real daily pressure: mixed-inventory orders, partial payments, unpaid follow-up, drafts, barcode use, branch-aware insights, supplier-linked stock, and wallet-based finance.",
        ],
    )
    project(
        doc,
        "AkananTV Desktop - Retail Operations App",
        "https://ashraf-swaidan.pages.dev/works/ak-system",
        [
            "Built a local-first desktop system for a TV/electronics retail workflow.",
            "Centralized inventory categories, sales, customer history, supplier workflows, receipt printing, protected private mode, bilingual Arabic/English UX, and backup/restore.",
        ],
    )
    project(
        doc,
        "Duwit - AI Goal Execution Platform",
        "https://ashraf-swaidan.pages.dev/works/duwit",
        [
            "Built an AI execution product that turns vague goals into phased plans, task-level coaching, quizzes, progress memory, and web/desktop delivery.",
            "Focused on helping users move from intention to completion instead of staying inside generic motivational chat.",
        ],
        live_url="https://duwit-45a37.web.app/",
    )
    project(
        doc,
        "Twodo - Collaborative Task App",
        "https://ashraf-swaidan.pages.dev/works/twodo",
        [
            "Built a low-friction task app around fast capture, simple project organization, invitations, collaboration, and a one-click UX mindset.",
        ],
        live_url="https://twodo.ashraf-swaidan-10.workers.dev/",
    )

    section(doc, "Experience")
    role(doc, "Papion Event Decorations / Print & Cut Workshop - Operations Lead and Internal Systems Developer", "Lebanon | 2020-2025")
    for item in [
        "Worked inside a real event decoration and print/cut business while studying Computer Science, eventually leading workshop operations and supporting customer-facing, production, and internal workflows.",
        "Identified operational friction across inventory, sales, customers, finance, suppliers, staff roles, and reporting, then designed software workflows around how the business actually worked.",
        "Built and iterated internal systems used in real daily operations, with a strong focus on usability, decision clarity, and reducing repeated manual work.",
    ]:
        bullet(doc, item)

    role(doc, "Independent Software Developer - Client Systems and Product Work", "Lebanon | 2022-2025")
    for item in [
        "Built business tools and personal products across web and desktop, including operational dashboards, local-first apps, AI-assisted planning tools, and collaborative task workflows.",
        "Worked across product thinking, UI/UX, frontend implementation, backend/data structure, desktop packaging, and deployment.",
        "Used AI development tools as an acceleration layer for planning, implementation, review, and iteration while keeping ownership of product direction and system behavior.",
    ]:
        bullet(doc, item)

    section(doc, "Skills")
    skill_line(doc, "Product and UX", "Workflow analysis, internal tools, operational systems, dashboard design, business process mapping, responsive interfaces, bilingual/RTL UX, iteration-driven product refinement")
    skill_line(doc, "Frontend", "React, TypeScript, Vite, React Router, TanStack Router, TanStack Query, Tailwind CSS, shadcn/Radix UI, Framer Motion, GSAP")
    skill_line(doc, "Backend, Data, and Desktop", "Firebase, Node.js, Express, MongoDB, LowDB, SQL/NoSQL, PHP, Django, Electron, IPC architecture, local-first persistence, backup/restore, desktop packaging")
    skill_line(doc, "Design, Media, and AI Workflow", "Photoshop, Illustrator, After Effects, Premiere Pro, print/cut production workflows, Cursor, Codex, prompt architecture, AI-assisted development and review")

    section(doc, "Education")
    role(doc, "Bachelor Degree in Computer Science", "Lebanese International University | Lebanon | Graduated late 2023")

    section(doc, "Additional Background")
    for item in [
        "Practical experience in customer-facing operations, workshop leadership, business analysis, production workflows, printing systems, machine-related software, and fabrication processes.",
        "Comfortable moving between business needs, interface design, and engineering execution.",
        "Strong preference for building useful systems over decorative demos or generic technical showcases.",
    ]:
        bullet(doc, item)

    doc.save(OUT)


if __name__ == "__main__":
    build()
